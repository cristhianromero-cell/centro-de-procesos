import streamlit as st
import pandas as pd
import re
import unicodedata
from io import BytesIO
from datetime import datetime
from openpyxl.styles import Font, PatternFill, Alignment

st.set_page_config(
    page_title="Centro de Procesos",
    page_icon="⚙️",
    layout="wide",
    initial_sidebar_state="expanded"
)

st.markdown(r"""
<style>
    .stApp {
        background: linear-gradient(180deg, #f5f8fc 0%, #eef3f8 100%);
    }
    .block-container {
        max-width: 1180px;
        padding-top: 2rem;
        padding-bottom: 3rem;
    }
    .hero {
        padding: 28px 30px;
        border-radius: 22px;
        background: linear-gradient(120deg, #0f2747 0%, #174f75 55%, #1f7a8c 100%);
        color: white;
        box-shadow: 0 18px 45px rgba(15,39,71,.18);
        margin-bottom: 22px;
    }
    .hero h1 { margin:0; font-size:2.05rem; font-weight:800; }
    .hero p { margin:8px 0 0 0; opacity:.88; font-size:1.02rem; }
    .process-card {
        border: 1px solid #dfe8f1;
        border-radius: 18px;
        padding: 20px;
        background: rgba(255,255,255,.94);
        min-height: 175px;
        box-shadow: 0 8px 26px rgba(31, 78, 120, .07);
    }
    .process-icon {font-size: 2rem; margin-bottom: 6px;}
    .process-card h3 {margin:4px 0 6px 0; color:#173451;}
    .process-card p {color:#607287; font-size:.92rem;}
    .badge {
        display:inline-block; padding:5px 9px; border-radius:999px;
        background:#e8f4f5; color:#11606a; font-size:.76rem; font-weight:700;
    }
    div[data-testid="stFileUploader"] {
        background:white; border:1px dashed #adc1d5; border-radius:16px; padding:8px 12px;
    }
    div[data-testid="stMetric"] {
        background:white; padding:14px; border-radius:15px; border:1px solid #e2eaf2;
    }
    .small-note {color:#6f8194; font-size:.86rem;}
    section[data-testid="stSidebar"] {background:#10243c;}
    section[data-testid="stSidebar"] * {color:#eef6fb;}
</style>
""", unsafe_allow_html=True)

def limpiar(texto):

    if pd.isna(texto):
        return ""

    texto = str(texto)

    texto = texto.replace("\xa0", " ")
    texto = texto.replace("\n", " ")
    texto = texto.replace("\r", " ")
    texto = texto.replace("“", '"')
    texto = texto.replace("”", '"')

    texto = re.sub(
        r"\s+",
        " ",
        texto
    )

    return texto.strip()


def sin_tildes(texto):

    texto = unicodedata.normalize(
        "NFD",
        texto
    )

    return "".join(
        c for c in texto
        if unicodedata.category(c) != "Mn"
    )


def normalizar(texto):

    return sin_tildes(
        limpiar(texto)
    ).upper()


def unicos(lista):

    resultado = []
    vistos = set()

    for x in lista:

        if x is None:
            continue

        x = str(x)

        x = re.sub(
            r"\s+",
            " ",
            x
        ).strip(" ,;/:-()")

        if not x:
            continue

        llave = normalizar(x)

        if llave not in vistos:
            vistos.add(llave)
            resultado.append(x)

    return resultado


FORMAS_FARMACEUTICAS = [

    (
        [
            "SOFTGEL CAPSULES",
            "SOFTGEL CAPSULE",
            "SOFT GEL CAPSULES",
            "SOFT GEL CAPSULE",
            "SOFT GELATIN CAPSULES",
            "SOFT GELATIN CAPSULE",
            "CAPSULA BLANDA",
            "CAPSULAS BLANDAS"
        ],
        "CÁPSULA BLANDA"
    ),

    (
        [
            "HARD GELATIN CAPSULE",
            "HARD GELATIN CAPSULES",
            "CAPSULA DURA",
            "CAPSULAS DURAS"
        ],
        "CÁPSULA DURA"
    ),

    (
        [
            "FILM COATED TABLETS",
            "FILM COATED TABLET",
            "COATED TABLETS",
            "COATED TABLET",
            "TABLETAS RECUBIERTAS",
            "TABLETA RECUBIERTA"
        ],
        "TABLETA RECUBIERTA"
    ),

    (
        [
            "DISPERSIBLE TABLET",
            "DISPERSIBLE TABLETS",
            "TABLETA DISPERSABLE",
            "TABLETAS DISPERSABLES"
        ],
        "TABLETA DISPERSABLE"
    ),

    (
        [
            "CHEWABLE TABLET",
            "CHEWABLE TABLETS",
            "TABLETA MASTICABLE",
            "TABLETAS MASTICABLES"
        ],
        "TABLETA MASTICABLE"
    ),

    (
        [
            "SUBLINGUAL TABLET",
            "SUBLINGUAL TABLETS",
            "TABLETA SUBLINGUAL",
            "TABLETAS SUBLINGUALES"
        ],
        "TABLETA SUBLINGUAL"
    ),

    (
        [
            "TABLETS",
            "TABLET",
            "TABLETAS",
            "TABLETA"
        ],
        "TABLETA"
    ),

    (
        [
            "CAPSULES",
            "CAPSULE",
            "CAPSULAS",
            "CAPSULA"
        ],
        "CÁPSULA"
    ),

    (
        [
            "POWDER FOR SOLUTION FOR INJECTION",
            "POLVO PARA SOLUCION INYECTABLE"
        ],
        "POLVO PARA SOLUCIÓN INYECTABLE"
    ),

    (
        [
            "SOLUTION FOR INJECTION",
            "INJECTION SOLUTION",
            "SOLUCION INYECTABLE"
        ],
        "SOLUCIÓN INYECTABLE"
    ),

    (
        [
            "SOLUTION FOR INFUSION",
            "SOLUCION PARA INFUSION"
        ],
        "SOLUCIÓN PARA INFUSIÓN"
    ),

    (
        [
            "ORAL SUSPENSION",
            "SUSPENSION ORAL"
        ],
        "SUSPENSIÓN ORAL"
    ),

    (
        [
            "ORAL SOLUTION",
            "SOLUCION ORAL"
        ],
        "SOLUCIÓN ORAL"
    ),

    (
        [
            "SYRUP",
            "JARABE"
        ],
        "JARABE"
    ),

    (
        [
            "CREAM",
            "CREMA"
        ],
        "CREMA"
    ),

    (
        [
            "OINTMENT",
            "UNGUENTO"
        ],
        "UNGÜENTO"
    ),

    (
        [
            "EYE DROPS",
            "GOTAS OFTALMICAS"
        ],
        "GOTAS OFTÁLMICAS"
    ),

    (
        [
            "ORAL DROPS",
            "GOTAS ORALES"
        ],
        "GOTAS ORALES"
    ),

    (
        [
            "SUPPOSITORY",
            "SUPPOSITORIES",
            "SUPOSITORIO",
            "SUPOSITORIOS"
        ],
        "SUPOSITORIO"
    ),

    (
        [
            "PATCH",
            "PATCHES",
            "PARCHE",
            "PARCHES"
        ],
        "PARCHE"
    ),

    (
        [
            "INHALER",
            "INHALADOR"
        ],
        "INHALADOR"
    ),

    (
        [
            "AEROSOL"
        ],
        "AEROSOL"
    ),

    (
        [
            "GEL"
        ],
        "GEL"
    )
]


def extraer_forma_farmaceutica(texto):

    texto_n = normalizar(texto)

    for variantes, resultado in FORMAS_FARMACEUTICAS:

        for variante in variantes:

            if normalizar(variante) in texto_n:
                return resultado

    return ""


def extraer_concentraciones(texto):

    texto_n = normalizar(texto)

    patrones = [

        r"\b\d+(?:[.,]\d+)?\s*MG\s*/\s*\d+(?:[.,]\d+)?\s*ML\b",

        r"\b\d+(?:[.,]\d+)?\s*MCG\s*/\s*\d+(?:[.,]\d+)?\s*ML\b",

        r"\b\d+(?:[.,]\d+)?\s*G\s*/\s*\d+(?:[.,]\d+)?\s*ML\b",

        r"\b\d+(?:[.,]\d+)?\s*MG\s*/\s*G\b",

        r"\b\d+(?:[.,]\d+)?\s*MG\b",

        r"\b\d+(?:[.,]\d+)?\s*MCG\b",

        r"\b\d+(?:[.,]\d+)?\s*UG\b",

        r"\b\d+(?:[.,]\d+)?\s*G\b",

        r"\b\d+(?:[.,]\d+)?\s*UI\b",

        r"\b\d+(?:[.,]\d+)?\s*IU\b",

        r"\b\d+(?:[.,]\d+)?\s*%\b"
    ]

    resultados = []

    for patron in patrones:

        encontrados = re.findall(
            patron,
            texto_n
        )

        resultados.extend(
            encontrados
        )

    return unicos(resultados)


def extraer_concentracion(texto):

    concentraciones = extraer_concentraciones(
        texto
    )

    return " + ".join(
        concentraciones
    )


PALABRAS_FORMA = [

    "SOFTGEL CAPSULES",
    "SOFTGEL CAPSULE",
    "SOFT GEL CAPSULE",
    "SOFT GEL CAPSULES",
    "SOFT GELATIN CAPSULE",
    "SOFT GELATIN CAPSULES",
    "FILM COATED TABLETS",
    "FILM COATED TABLET",
    "COATED TABLETS",
    "COATED TABLET",
    "TABLETS",
    "TABLET",
    "TABLETAS",
    "TABLETA",
    "CAPSULES",
    "CAPSULE",
    "CAPSULAS",
    "CAPSULA",
    "INJECTION",
    "INYECTABLE",
    "SOLUTION",
    "SOLUCION",
    "SUSPENSION",
    "SYRUP",
    "JARABE",
    "CREAM",
    "CREMA",
    "GEL",
    "VIAL",
    "AMPOULE",
    "AMPOLLA",
    "POWDER",
    "POLVO"
]


PALABRAS_ADMINISTRATIVAS = [

    "LOTE",
    "LOT",
    "BATCH",
    "MFG",
    "MFG DATE",
    "EXP",
    "EXP DATE",
    "EXPIRY",
    "DATE",
    "REGISTRO",
    "SANITARIO",
    "SANITARIA",
    "USO",
    "PAGO",
    "CREDITO",
    "FACTURA",
    "INVOICE",
    "EXPEDIENTE",
    "SUCE",
    "MEDICAMENTO",
    "CONSUMO HUMANO",
    "PACKING",
    "PACK SIZE"
]


def limpiar_molecula(texto):

    resultado = normalizar(texto)

    resultado = re.sub(
        r"\b\d+(?:[.,]\d+)?\s*(?:MG|MCG|UG|G|ML|UI|IU|%)"
        r"(?:\s*/\s*\d+(?:[.,]\d+)?\s*(?:ML|G))?\b",
        " ",
        resultado
    )

    for palabra in PALABRAS_FORMA:

        resultado = re.sub(
            r"\b"
            + re.escape(normalizar(palabra))
            + r"\b",
            " ",
            resultado
        )

    resultado = re.sub(
        r"\b(?:EACH|CONTAINS|CONTAINING|COMPOSITION|COMPOSICION)\b",
        " ",
        resultado
    )

    resultado = re.sub(
        r"\s+",
        " ",
        resultado
    )

    resultado = resultado.strip(
        " ,;:/()-"
    )

    return resultado


def es_segmento_administrativo(texto):

    texto_n = normalizar(texto)

    for palabra in PALABRAS_ADMINISTRATIVAS:

        if normalizar(palabra) in texto_n:
            return True

    return False


def extraer_bloques_parentesis(texto):

    return re.findall(
        r"\(([^()]*)\)",
        texto
    )


def puntuar_bloque_molecula(bloque):

    bloque_n = normalizar(bloque)

    puntos = 0

    concentraciones = extraer_concentraciones(
        bloque
    )

    if concentraciones:
        puntos += 3

    if "+" in bloque:
        puntos += 4

    if len(concentraciones) >= 2:
        puntos += 3

    if any(
        normalizar(f) in bloque_n
        for f in PALABRAS_FORMA
    ):
        puntos += 1

    if es_segmento_administrativo(
        bloque
    ):
        puntos -= 5

    return puntos


def extraer_bloque_composicion(texto):

    candidatos = []

    parentesis = extraer_bloques_parentesis(
        texto
    )

    for bloque in parentesis:

        puntos = puntuar_bloque_molecula(
            bloque
        )

        if puntos > 0:

            candidatos.append(
                (
                    puntos,
                    bloque
                )
            )

    segmentos = re.split(
        r"//",
        texto
    )

    for segmento in segmentos:

        puntos = puntuar_bloque_molecula(
            segmento
        )

        if puntos > 0:

            candidatos.append(
                (
                    puntos,
                    segmento
                )
            )

    if not candidatos:
        return ""

    candidatos.sort(
        key=lambda x: x[0],
        reverse=True
    )

    return candidatos[0][1]


def extraer_moleculas_combinadas(texto):

    bloque = extraer_bloque_composicion(
        texto
    )

    if not bloque:
        return ""

    if "+" in bloque:

        partes = re.split(
            r"\s*\+\s*",
            bloque
        )

        moleculas = []

        for parte in partes:

            limpia = limpiar_molecula(
                parte
            )

            if limpia:
                moleculas.append(
                    limpia
                )

        moleculas = unicos(
            moleculas
        )

        if len(moleculas) >= 2:

            return " + ".join(
                moleculas
            )

    return ""


def extraer_molecula_simple(texto):

    bloque = extraer_bloque_composicion(
        texto
    )

    if not bloque:
        return ""

    molecula = limpiar_molecula(
        bloque
    )

    if not molecula:
        return ""

    if len(molecula) > 100:
        return ""

    return molecula


def extraer_molecula(texto):

    combinacion = extraer_moleculas_combinadas(
        texto
    )

    if combinacion:
        return combinacion

    return extraer_molecula_simple(
        texto
    )


def extraer_marca(texto):

    texto = limpiar(
        texto
    )

    primera_parte = texto.split(
        "//"
    )[0]

    partes = [
        p.strip()
        for p in primera_parte.split(",")
        if p.strip()
    ]

    if len(partes) >= 2:

        posible_marca = partes[1]

        posible_n = normalizar(
            posible_marca
        )

        if posible_n not in [
            "S/M",
            "SM",
            "S M",
            "SIN MARCA",
            "NA",
            "N/A"
        ]:

            return posible_marca.upper()

    if partes:

        primera = normalizar(
            partes[0]
        )

        for forma in PALABRAS_FORMA:

            primera = re.sub(
                r"\b"
                + re.escape(
                    normalizar(forma)
                )
                + r"\b",
                " ",
                primera
            )

        primera = re.sub(
            r"\b\d+(?:[.,]\d+)?\s*(?:MG|MCG|UG|G|ML|UI|IU|%)\b",
            " ",
            primera
        )

        primera = re.sub(
            r"\s+",
            " ",
            primera
        ).strip()

        if len(primera) <= 60:
            return primera

    return ""


def extraer_presentacion(texto):

    texto_n = normalizar(
        texto
    )

    patrones = [

        r"\bCAJA\s+(?:X|POR|CON)\s+\d+[^/;,]*",

        r"\bBOX\s+(?:OF|X)?\s*\d+[^/;,]*",

        r"\bFRASCO\s+(?:X|POR|CON)?\s*\d+(?:[.,]\d+)?\s*(?:ML|G|MG)?[^/;,]*",

        r"\bBOTTLE\s+(?:OF|X)?\s*\d+(?:[.,]\d+)?\s*(?:ML|G|MG)?[^/;,]*",

        r"\bBLISTER\s+(?:X|POR|CON)?\s*\d+[^/;,]*",

        r"\b\d+\s*BLISTERS?\s*(?:X|OF)?\s*\d+[^/;,]*",

        r"\b\d+\s*X\s*\d+\s*(?:TABLETS?|TABLETAS?|CAPSULES?|CAPSULAS?|AMPOLLAS?|VIALS?|SACHETS?)\b",

        r"\bTUBO\s+(?:X|DE)?\s*\d+(?:[.,]\d+)?\s*(?:G|ML)\b",

        r"\bTUBE\s+(?:X|OF)?\s*\d+(?:[.,]\d+)?\s*(?:G|ML)\b",

        r"\bVIAL\s+(?:X|DE)?\s*\d+(?:[.,]\d+)?\s*ML\b",

        r"\bAMPOLLA\s+(?:X|DE)?\s*\d+(?:[.,]\d+)?\s*ML\b"
    ]

    encontrados = []

    for patron in patrones:

        encontrados.extend(
            re.findall(
                patron,
                texto_n
            )
        )

    encontrados = unicos(
        encontrados
    )

    if encontrados:
        return " | ".join(
            encontrados[:3]
        )

    return ""


def crear_marca_concentracion(
    marca,
    concentracion
):

    if marca and concentracion:
        return (
            f"{marca} {concentracion}"
        )

    if marca:
        return marca

    if concentracion:
        return concentracion

    return ""


def procesar_descripcion(texto):

    original = (
        ""
        if pd.isna(texto)
        else str(texto)
    )

    limpio = limpiar(
        original
    )

    if not limpio:

        return {

            "DESCRIPCION_ORIGINAL":
                original,

            "MARCA_Y_CONCENTRACION":
                "",

            "MARCA":
                "",

            "CONCENTRACION":
                "",

            "MOLECULA":
                "",

            "FORMA_PRESENTACION":
                "",

            "FORMA_FARMACEUTICA":
                ""
        }

    marca = extraer_marca(
        limpio
    )

    concentracion = extraer_concentracion(
        limpio
    )

    molecula = extraer_molecula(
        limpio
    )

    presentacion = extraer_presentacion(
        limpio
    )

    forma = extraer_forma_farmaceutica(
        limpio
    )

    marca_concentracion = (
        crear_marca_concentracion(
            marca,
            concentracion
        )
    )

    return {

        "DESCRIPCION_ORIGINAL":
            original,

        "MARCA_Y_CONCENTRACION":
            marca_concentracion,

        "MARCA":
            marca,

        "CONCENTRACION":
            concentracion,

        "MOLECULA":
            molecula,

        "FORMA_PRESENTACION":
            presentacion,

        "FORMA_FARMACEUTICA":
            forma
    }



def preparar_resultado(archivo_subido, hoja):
    df = pd.read_excel(archivo_subido, sheet_name=hoja)
    df.columns = (
        df.columns.astype(str)
        .str.replace("\n", " ", regex=False)
        .str.replace("\r", " ", regex=False)
        .str.replace(r"\s+", " ", regex=True)
        .str.strip()
    )

    columna = next(
        (c for c in df.columns if c.strip().lower() in ["descripcion", "descripción"]),
        None
    )
    if columna is None:
        raise ValueError("No encontré una columna llamada Descripcion o Descripción en la hoja seleccionada.")

    resultado = pd.DataFrame([procesar_descripcion(x) for x in df[columna]])

    salida = BytesIO()
    with pd.ExcelWriter(salida, engine="openpyxl") as writer:
        resultado.to_excel(writer, index=False, sheet_name="PRODUCTOS_LIMPIOS")
        ws = writer.book["PRODUCTOS_LIMPIOS"]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for celda in ws[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="1F4E78")
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        anchos = {"A": 90, "B": 40, "C": 30, "D": 35, "E": 55, "F": 50, "G": 35}
        for letra, ancho in anchos.items():
            ws.column_dimensions[letra].width = ancho

        for fila in ws.iter_rows(min_row=2):
            for celda in fila:
                celda.alignment = Alignment(vertical="top", wrap_text=True)

    salida.seek(0)
    return resultado, salida.getvalue(), len(df)


def resumen_deteccion(resultado):
    campos = ["MARCA", "CONCENTRACION", "MOLECULA", "FORMA_PRESENTACION", "FORMA_FARMACEUTICA"]
    return {campo: int(resultado[campo].fillna("").ne("").sum()) for campo in campos}


if "vista" not in st.session_state:
    st.session_state.vista = "inicio"
if "resultado_excel" not in st.session_state:
    st.session_state.resultado_excel = None
if "resultado_df" not in st.session_state:
    st.session_state.resultado_df = None
if "nombre_salida" not in st.session_state:
    st.session_state.nombre_salida = None

with st.sidebar:
    st.markdown("## ⚙️ Centro de procesos")
    st.caption("Automatización de archivos")
    st.divider()
    if st.button("🏠 Inicio", use_container_width=True):
        st.session_state.vista = "inicio"
    if st.button("💊 Limpieza farmacéutica", use_container_width=True):
        st.session_state.vista = "farmaceutica"
    st.divider()
    st.caption("Nuevos procesos podrán agregarse aquí sin cambiar la estructura principal.")

st.markdown("""
<div class="hero">
    <h1>Centro de Procesos Automatizados</h1>
    <p>Sube tus archivos, selecciona el proceso y descarga el resultado listo para trabajar.</p>
</div>
""", unsafe_allow_html=True)

if st.session_state.vista == "inicio":
    st.subheader("Procesos disponibles")
    c1, c2, c3 = st.columns(3)
    with c1:
        st.markdown("""
        <div class="process-card">
            <div class="process-icon">💊</div>
            <span class="badge">ACTIVO</span>
            <h3>Limpieza farmacéutica</h3>
            <p>Extrae marca, concentración, molécula, presentación y forma farmacéutica desde la descripción.</p>
        </div>
        """, unsafe_allow_html=True)
        if st.button("Abrir proceso", key="abrir_farma", type="primary", use_container_width=True):
            st.session_state.vista = "farmaceutica"
            st.rerun()
    with c2:
        st.markdown("""
        <div class="process-card">
            <div class="process-icon">📊</div>
            <span class="badge">PRÓXIMAMENTE</span>
            <h3>Proceso 2</h3>
            <p>Espacio preparado para tu siguiente automatización de Excel o análisis de datos.</p>
        </div>
        """, unsafe_allow_html=True)
        st.button("Aún no disponible", disabled=True, use_container_width=True, key="p2")
    with c3:
        st.markdown("""
        <div class="process-card">
            <div class="process-icon">🧩</div>
            <span class="badge">PRÓXIMAMENTE</span>
            <h3>Proceso 3</h3>
            <p>Podemos agregar más módulos y cada uno tendrá sus propios controles y descargas.</p>
        </div>
        """, unsafe_allow_html=True)
        st.button("Aún no disponible", disabled=True, use_container_width=True, key="p3")

else:
    top1, top2 = st.columns([5,1])
    with top1:
        st.subheader("💊 Limpieza de Descripción")
        st.markdown('<div class="small-note">Carga un archivo Excel, selecciona la hoja y ejecuta el proceso.</div>', unsafe_allow_html=True)
    with top2:
        if st.button("← Volver", use_container_width=True):
            st.session_state.vista = "inicio"
            st.rerun()

    st.markdown("### 1. Cargar archivo")
    archivo = st.file_uploader("Arrastra o selecciona tu archivo Excel", type=["xlsx", "xls"], label_visibility="collapsed")

    if archivo is not None:
        try:
            excel = pd.ExcelFile(archivo)
            st.success(f"Archivo cargado correctamente: **{archivo.name}**")
            st.markdown("### 2. Elegir hoja")
            hoja = st.selectbox("Hoja a procesar", excel.sheet_names)

            with st.expander("Vista previa de la hoja", expanded=False):
                archivo.seek(0)
                preview = pd.read_excel(archivo, sheet_name=hoja, nrows=8)
                st.dataframe(preview, use_container_width=True, hide_index=True)

            st.markdown("### 3. Ejecutar")
            ejecutar = st.button("▶ Procesar archivo", type="primary", use_container_width=True)

            if ejecutar:
                archivo.seek(0)
                with st.spinner("Procesando registros..."):
                    resultado, bytes_excel, total = preparar_resultado(archivo, hoja)
                    st.session_state.resultado_df = resultado
                    st.session_state.resultado_excel = bytes_excel
                    fecha = datetime.now().strftime("%Y%m%d_%H%M%S")
                    hoja_segura = re.sub(r'[\/*?:\[\]]', '_', hoja)
                    st.session_state.nombre_salida = f"{hoja_segura}_PRODUCTOS_FARMACEUTICOS_{fecha}.xlsx"
                st.success("Proceso terminado correctamente.")

            if st.session_state.resultado_df is not None:
                resultado = st.session_state.resultado_df
                deteccion = resumen_deteccion(resultado)
                st.markdown("### Resultado")
                m1, m2, m3, m4, m5 = st.columns(5)
                total_r = len(resultado)
                for col, campo, titulo in [
                    (m1,"MARCA","Marcas"),
                    (m2,"CONCENTRACION","Concentraciones"),
                    (m3,"MOLECULA","Moléculas"),
                    (m4,"FORMA_PRESENTACION","Presentaciones"),
                    (m5,"FORMA_FARMACEUTICA","Formas farmacéuticas")
                ]:
                    with col:
                        cant = deteccion[campo]
                        pct = (cant/total_r*100) if total_r else 0
                        st.metric(titulo, f"{cant:,}", f"{pct:.1f}% detectado")

                with st.expander("Ver primeras filas del resultado", expanded=True):
                    st.dataframe(resultado.head(50), use_container_width=True, hide_index=True)

                st.download_button(
                    "⬇ Descargar Excel procesado",
                    data=st.session_state.resultado_excel,
                    file_name=st.session_state.nombre_salida,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    type="primary",
                    use_container_width=True
                )
        except Exception as e:
            st.error(f"No se pudo procesar el archivo: {e}")
